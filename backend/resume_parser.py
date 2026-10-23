import re
import spacy
from spacy.matcher import PhraseMatcher
from pypdf import PdfReader
from docx import Document
from skills import SKILLS, ROLES

nlp = spacy.load("en_core_web_sm")


def build_matcher(table):
    m = PhraseMatcher(nlp.vocab, attr="LOWER")
    for canonical, aliases in table.items():
        m.add(canonical, [nlp.make_doc(n) for n in [canonical] + aliases])
    return m


skill_matcher = build_matcher(SKILLS)
role_matcher = build_matcher(ROLES)

# Multi-word skills are also searched as plain text, to survive PDFs that
# glue words together (for example "theMERN Stack")
MULTIWORD = []
for _canonical, _aliases in SKILLS.items():
    for _name in [_canonical] + _aliases:
        if " " in _name:
            MULTIWORD.append((_name.lower(), _canonical))


def find_glued(lower_text, already):
    extra = []
    for phrase, canonical in MULTIWORD:
        if canonical not in already and canonical not in extra and phrase in lower_text:
            extra.append(canonical)
    return extra
# These skills are also ordinary words, so they only count when capitalised
CASE_SENSITIVE = {"Excel", "Swift", "Render", "C"}

DEGREES = [  # highest first
    ("PhD", r"\b(ph\.?d|doctorate)\b"),
    ("Master's degree", r"\b(m\.?\s?tech|m\.?\s?sc|mba|mca|masters?|master['’]s|m\.s\.|m\.e\.?)(?![a-z])"),
    ("Bachelor's degree", r"\b(b\.?\s?tech|b\.?\s?sc|bca|b\.?\s?com|bachelors?|bachelor['’]s|b\.e\.?|b\.a\.|b\.s\.)(?![a-z])"),
    ("Diploma", r"\bdiploma\b"),
]

FIELDS = [
    "computer science", "information technology", "data science",
    "artificial intelligence", "electronics", "electrical", "mechanical",
    "civil", "business administration", "statistics", "mathematics",
    "physics", "commerce",
]


def extract_text(file_storage, ext):
    """Reads plain text out of an uploaded PDF or DOCX."""
    if ext == "pdf":
        reader = PdfReader(file_storage.stream)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    doc = Document(file_storage.stream)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def find_matches(matcher, doc, case_sensitive=()):
    """Names found by a matcher, in the order they first appear."""
    found = {}
    for match_id, start, end in matcher(doc):
        name = nlp.vocab.strings[match_id]
        span = doc[start:end]
        # Ambiguous single words only count when they start with a capital
        if name in case_sensitive and len(span) == 1 and span.text[0].islower():
            continue
        # A lone "C" inside "C++" or "C#" is not the C language
        if name == "C":
            after = doc.text[span.end_char:span.end_char + 1]
            if after in ("+", "#"):
                continue
        found.setdefault(name, start)
    return sorted(found, key=found.get)


def find_education(text):
    lower = text.lower()
    for label, pattern in DEGREES:
        if re.search(pattern, lower):
            return label
    return ""


def find_field(text):
    lower = text.lower()
    for field in FIELDS:
        if field in lower:
            return field.title()
    return ""


def find_experience_years(text):
    lower = text.lower()
    nums = re.findall(
        r"(\d{1,2})\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:\w+\s+){0,2}experience",
        lower,
    )
    nums += re.findall(
        r"experience\s*(?:of|:)?\s*(\d{1,2})\+?\s*(?:years?|yrs?)",
        lower,
    )
    return max((int(n) for n in nums), default=0)


def find_certifications(text):
    certs = []
    for line in text.splitlines():
        line = line.strip(" \t•-–*·")
        if 5 < len(line) <= 80 and re.search(r"certif", line, re.I):
            if line not in certs:
                certs.append(line)
    return certs[:10]


def parse_resume(text):
    flat = re.sub(r"\s+", " ", text)   # line breaks inside a phrase break matching
    doc = nlp(flat)
    skills = find_matches(skill_matcher, doc, CASE_SENSITIVE)
    skills += find_glued(flat.lower(), set(skills))
    return {
        "skills": skills,
        "roles": find_matches(role_matcher, doc),
        "education": find_education(flat),
        "field_of_study": find_field(flat),
        "experience_years": find_experience_years(flat),
        "certifications": find_certifications(text),
    }